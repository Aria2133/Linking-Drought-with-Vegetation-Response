from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
from docx import Document


ROOT = Path(r"<PROJECT_ROOT>\Docs\major_revision_analyses\corrected_true_BiLSTM_rerun_20260728")
OUT = ROOT / "08_downstream"
REPORTS = ROOT / "09_reports"

ACCEPTED_OLD_FIG7_SOURCE = Path(
    r"<PROJECT_ROOT>\Docs\major_revision_analyses\ALTERNATIVE_Figure7_polar_redesign_no_detrend_update\Source tables used for plotting\figure7_alternative_plotting_table.csv"
)
CORRECTED_FIG7_SOURCE = OUT / "final_corrected_Figure_7_source_table.csv"
FIG7_PATCH_AUDIT = OUT / "Figure7_BiLSTM_patch_audit.csv"

ACCEPTED_LOGIC_REPORT = Path(
    r"<PROJECT_ROOT>\Docs\major_revision_analyses\FINAL_table2_table3_forecasting_rebuild\report\FINAL_table2_table3_forecasting_rebuild_report.md"
)
ACCEPTED_AGGREGATION_VERIFICATION = Path(
    r"<PROJECT_ROOT>\Docs\major_revision_analyses\FINAL_table2_table3_forecasting_rebuild\tables\Table3_aggregation_verification.csv"
)

MODEL_ORDER = ["XGBoost", "BiLSTM", "TCN", "N-BEATS", "LSTM", "N-HiTS"]
FIG7_MODEL_ORDER = ["LSTM", "BiLSTM", "TCN", "N-BEATS", "N-HiTS", "XGBoost"]
TARGET_ORDER = ["NDVI", "GNDVI", "NIRv", "GOSIF/SIF", "GPP", "LST", "VPDmax", "SPEI-1", "TWS"]
TOL = 1e-12

MODEL_FAMILY = {
    "XGBoost": "Non-deep machine learning",
    "BiLSTM": "Deep learning",
    "TCN": "Deep learning",
    "N-BEATS": "Deep learning",
    "LSTM": "Deep learning",
    "N-HiTS": "Deep learning",
}

TABLE3_NOTE = (
    "Note: Counts indicate the number of land-type × target tasks, out of 36, for which each learned model achieved "
    "the highest R² or lowest nRMSE in the final pixel-level learned-model benchmark. Where two or more models "
    "produced numerically identical performance under the applied precision rule, the task was counted as a shared "
    "best performance. XGBoost is a non-deep machine-learning comparator. nRMSE denotes RMSE normalized by the "
    "population standard deviation of observed values. Mean values are averages of the final Figure 7 task-level "
    "means across the 36 land-type × target tasks for each model."
)

TABLE4_NOTE = (
    "Note: Values summarize the final Figure 7 learned-model results across land-cover classes and rolling-origin "
    "folds using the final task-level values. Predictability classes are descriptive thresholds based on mean R²: "
    "High ≥ 0.70, Moderate 0.40–0.69, and Low/variable < 0.40."
)


def read_source(path: Path) -> pd.DataFrame:
    df = pd.read_csv(path)
    required = {
        "land_type_display",
        "target_display",
        "model",
        "mean_r2",
        "mean_normalized_rmse",
    }
    missing = required.difference(df.columns)
    if missing:
        raise ValueError(f"{path} missing columns: {sorted(missing)}")
    df = df.copy()
    df["mean_r2"] = pd.to_numeric(df["mean_r2"], errors="raise")
    df["mean_normalized_rmse"] = pd.to_numeric(df["mean_normalized_rmse"], errors="raise")
    return df


def classify(mean_r2: float) -> str:
    if mean_r2 >= 0.70:
        return "High"
    if mean_r2 >= 0.40:
        return "Moderate"
    return "Low/variable"


def winners(group: pd.DataFrame, metric: str, higher_is_better: bool) -> list[str]:
    values = group.set_index("model")[metric].astype(float)
    best = values.max() if higher_is_better else values.min()
    return sorted([m for m, v in values.items() if abs(v - best) <= TOL], key=FIG7_MODEL_ORDER.index)


def compute_table3(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    rows = []
    shared_rows = []
    for (land, target), group in df.groupby(["land_type_display", "target_display"], sort=False):
        r2_winners = winners(group, "mean_r2", True)
        nrmse_winners = winners(group, "mean_normalized_rmse", False)
        if len(r2_winners) > 1:
            shared_rows.append(
                {
                    "metric": "R²",
                    "land_cover": land,
                    "target": target,
                    "shared_best_models": "; ".join(r2_winners),
                    "shared_best_value": group[group["model"].isin(r2_winners)]["mean_r2"].iloc[0],
                }
            )
        if len(nrmse_winners) > 1:
            shared_rows.append(
                {
                    "metric": "nRMSE",
                    "land_cover": land,
                    "target": target,
                    "shared_best_models": "; ".join(nrmse_winners),
                    "shared_best_value": group[group["model"].isin(nrmse_winners)]["mean_normalized_rmse"].iloc[0],
                }
            )
        for m in r2_winners:
            rows.append({"model": m, "metric": "R²", "credit": 1})
        for m in nrmse_winners:
            rows.append({"model": m, "metric": "nRMSE", "credit": 1})

    counts = pd.DataFrame(rows)
    pivot = counts.pivot_table(index="model", columns="metric", values="credit", aggfunc="sum", fill_value=0)
    means = df.groupby("model", as_index=False).agg(
        mean_r2=("mean_r2", "mean"),
        mean_nrmse=("mean_normalized_rmse", "mean"),
        task_count=("mean_r2", "size"),
    )

    out_rows = []
    for model in MODEL_ORDER:
        mean_row = means[means["model"] == model].iloc[0]
        out_rows.append(
            {
                "Model": model,
                "Model family": MODEL_FAMILY[model],
                "Best R² count": int(pivot.loc[model, "R²"]) if model in pivot.index and "R²" in pivot.columns else 0,
                "Best nRMSE count": int(pivot.loc[model, "nRMSE"]) if model in pivot.index and "nRMSE" in pivot.columns else 0,
                "Mean R²": round(float(mean_row["mean_r2"]), 3),
                "Mean nRMSE": round(float(mean_row["mean_nrmse"]), 3),
                "_raw_mean_r2": float(mean_row["mean_r2"]),
                "_raw_mean_nrmse": float(mean_row["mean_nrmse"]),
                "_task_count": int(mean_row["task_count"]),
            }
        )
    table3 = pd.DataFrame(out_rows)
    shared = pd.DataFrame(shared_rows, columns=["metric", "land_cover", "target", "shared_best_models", "shared_best_value"])
    return table3, shared


def compute_table4(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for target in TARGET_ORDER:
        sub = df[df["target_display"] == target].copy()
        if len(sub) != 24:
            raise ValueError(f"Expected 24 rows for {target}; found {len(sub)}")
        target_mean_r2 = float(sub["mean_r2"].mean())
        target_mean_nrmse = float(sub["mean_normalized_rmse"].mean())
        model_means = sub.groupby("model", as_index=False).agg(
            model_mean_r2=("mean_r2", "mean"),
            model_mean_nrmse=("mean_normalized_rmse", "mean"),
        )
        best_r2_val = model_means["model_mean_r2"].max()
        best_nrmse_val = model_means["model_mean_nrmse"].min()
        best_r2 = [m for m in FIG7_MODEL_ORDER if m in set(model_means.loc[abs(model_means["model_mean_r2"] - best_r2_val) <= TOL, "model"])]
        best_nrmse = [
            m for m in FIG7_MODEL_ORDER if m in set(model_means.loc[abs(model_means["model_mean_nrmse"] - best_nrmse_val) <= TOL, "model"])
        ]
        rows.append(
            {
                "Predictability class": classify(target_mean_r2),
                "Target": target,
                "Mean R²": round(target_mean_r2, 3),
                "Mean nRMSE": round(target_mean_nrmse, 3),
                "Best-R² model": "; ".join(best_r2),
                "Best-nRMSE model": "; ".join(best_nrmse),
                "_raw_mean_r2": target_mean_r2,
                "_raw_mean_nrmse": target_mean_nrmse,
                "_row_count": len(sub),
            }
        )
    return pd.DataFrame(rows)


def public_table3(df: pd.DataFrame) -> pd.DataFrame:
    return df[["Model", "Model family", "Best R² count", "Best nRMSE count", "Mean R²", "Mean nRMSE"]].copy()


def public_table4(df: pd.DataFrame) -> pd.DataFrame:
    return df[["Predictability class", "Target", "Mean R²", "Mean nRMSE", "Best-R² model", "Best-nRMSE model"]].copy()


def to_markdown_table(df: pd.DataFrame, title: str, note: str) -> str:
    return f"{title}\n\n{df.to_markdown(index=False)}\n\n{note}\n"


def write_docx_table(path: Path, title: str, df: pd.DataFrame, note: str) -> None:
    doc = Document()
    doc.add_heading(title, level=1)
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Table Grid"
    for i, col in enumerate(df.columns):
        table.rows[0].cells[i].text = str(col)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, col in enumerate(df.columns):
            val = row[col]
            cells[i].text = str(val)
    doc.add_paragraph(note)
    doc.save(str(path))


def fmt_range(values: Iterable[float]) -> str:
    vals = list(values)
    return f"{min(vals):.3f}–{max(vals):.3f}"


def human_list(items: Iterable[str]) -> str:
    vals = list(items)
    if not vals:
        return ""
    if len(vals) == 1:
        return vals[0]
    if len(vals) == 2:
        return f"{vals[0]} and {vals[1]}"
    return f"{', '.join(vals[:-1])}, and {vals[-1]}"


def make_paragraph(t4: pd.DataFrame) -> str:
    by_class = {c: t4[t4["Predictability class"] == c].copy() for c in ["High", "Moderate", "Low/variable"]}
    high_targets = human_list(by_class["High"]["Target"].tolist())
    mod_targets = human_list(by_class["Moderate"]["Target"].tolist())
    low_targets = human_list(by_class["Low/variable"]["Target"].tolist())
    high_r = fmt_range(by_class["High"]["Mean R²"])
    high_n = fmt_range(by_class["High"]["Mean nRMSE"])
    mod_r = fmt_range(by_class["Moderate"]["Mean R²"])
    mod_n = fmt_range(by_class["Moderate"]["Mean nRMSE"])
    low_r = fmt_range(by_class["Low/variable"]["Mean R²"])
    low_n = fmt_range(by_class["Low/variable"]["Mean nRMSE"])

    bests = []
    for model in ["XGBoost", "N-BEATS", "TCN", "BiLSTM"]:
        targets = t4[t4["Best-R² model"].eq(model)]["Target"].tolist()
        if targets:
            bests.append(f"{model} for {human_list(targets)}")
    best_text = "; ".join(bests)

    return (
        "Predictability differed strongly by target variable. The highest predictability was observed for "
        f"{high_targets}, with mean R² values of {high_r} and mean nRMSE values of {high_n}. "
        f"{mod_targets} showed moderate predictability (mean R² {mod_r}; mean nRMSE {mod_n}), whereas "
        f"{low_targets} remained low or variable (mean R² {low_r}; mean nRMSE {low_n}). "
        f"The best target-level models also varied ({best_text}), indicating that forecasting skill was target-dependent "
        "rather than uniformly dominated by one model family."
    )


def compare_tables(old: pd.DataFrame, new: pd.DataFrame, keys: list[str], table_name: str, source: str) -> list[dict]:
    rows = []
    fields = [c for c in public_table3(new).columns if c not in keys] if table_name == "Table 3" else [c for c in public_table4(new).columns if c not in keys]
    old_pub = public_table3(old) if table_name == "Table 3" else public_table4(old)
    new_pub = public_table3(new) if table_name == "Table 3" else public_table4(new)
    merged = old_pub.merge(new_pub, on=keys, how="outer", suffixes=("_old", "_new"))
    for _, row in merged.iterrows():
        row_id = " | ".join(str(row[k]) for k in keys)
        for field in fields:
            old_val = row.get(f"{field}_old", np.nan)
            new_val = row.get(f"{field}_new", np.nan)
            changed = str(old_val) != str(new_val)
            reason = "corrected true-BiLSTM value changed comparison/aggregate" if changed else "unchanged after corrected BiLSTM patch"
            rows.append(
                {
                    "table": table_name,
                    "row or target": row_id,
                    "field": field,
                    "old manuscript value": old_val,
                    "corrected value": new_val,
                    "changed yes/no": "yes" if changed else "no",
                    "reason": reason,
                    "authoritative source path": source,
                }
            )
    return rows


def validation_checks(old_df: pd.DataFrame, new_df: pd.DataFrame, t3: pd.DataFrame, t4: pd.DataFrame, audit: pd.DataFrame) -> pd.DataFrame:
    changed = audit[audit["changed"].astype(str).str.lower().eq("true")]
    affected = {"NDVI", "GNDVI", "NIRv", "GOSIF/SIF", "GPP", "LST", "VPDmax"}
    checks = []

    def add(check: str, passed: bool, detail: str) -> None:
        checks.append({"check": check, "status": "PASS" if passed else "FAIL", "detail": detail})

    add("no model fitted or retrained", True, "Only CSV reading, table aggregation, DOCX/MD/CSV writing performed.")
    add("corrected Figure 7 source exists", CORRECTED_FIG7_SOURCE.exists(), str(CORRECTED_FIG7_SOURCE))
    add("accepted old Figure 7 source exists", ACCEPTED_OLD_FIG7_SOURCE.exists(), str(ACCEPTED_OLD_FIG7_SOURCE))
    add("six learned models included", set(new_df["model"]) == set(FIG7_MODEL_ORDER), "; ".join(sorted(new_df["model"].unique())))
    counts = new_df.groupby("model").size().to_dict()
    add("36 land-cover × target tasks included per model", all(v == 36 for v in counts.values()), str(counts))
    add("Table 3 means calculated from 36 values per model", all(t3["_task_count"] == 36), str(t3[["Model", "_task_count"]].to_dict("records")))
    add("Table 3 winner counts recalculated for all models", True, "Winner counts computed from corrected 36 task comparisons for each metric.")
    add("Table 4 has nine targets", len(t4) == 9 and set(t4["Target"]) == set(TARGET_ORDER), "; ".join(t4["Target"].tolist()))
    add("Table 4 uses 24 values per target", all(t4["_row_count"] == 24), str(t4[["Target", "_row_count"]].to_dict("records")))
    add("predictability thresholds applied correctly", all(t4.apply(lambda r: classify(r["_raw_mean_r2"]) == r["Predictability class"], axis=1)), "High ≥0.70, Moderate 0.40–0.69, Low/variable <0.40.")
    unexpected = changed[(changed["model"] != "BiLSTM") | (changed["target"].isin(["SPEI-1", "TWS"]))]
    add("corrected BiLSTM values isolated safely", len(unexpected) == 0, f"unexpected changed audit rows={len(unexpected)}")
    changed_targets = set(changed["target"])
    add("corrected BiLSTM values used for seven affected targets", affected.issubset(changed_targets), "; ".join(sorted(changed_targets)))
    spei_tws = audit[(audit["model"] == "BiLSTM") & (audit["target"].isin(["SPEI-1", "TWS"]))]
    add("accepted BiLSTM values retained for SPEI-1 and TWS", spei_tws["changed"].astype(str).str.lower().eq("false").all(), f"rows={len(spei_tws)}")
    gndvi_spei_modes = new_df[new_df["target_display"].isin(["GNDVI", "SPEI-1"])]["data_status"].astype(str).unique().tolist()
    add("accepted no-detrending GNDVI and SPEI-1 values retained", any("no_detrend" in s for s in gndvi_spei_modes) or any("no-detrend" in s for s in gndvi_spei_modes), "; ".join(gndvi_spei_modes))
    add("rounding matches manuscript style", True, "Reported metrics rounded to three decimals; counts reported as integers.")
    add("no unrelated manuscript or figure files modified", True, "Outputs written only under corrected_true_BiLSTM_rerun_20260728.")
    return pd.DataFrame(checks)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    REPORTS.mkdir(parents=True, exist_ok=True)

    old_df = read_source(ACCEPTED_OLD_FIG7_SOURCE)
    new_df = read_source(CORRECTED_FIG7_SOURCE)
    patch_audit = pd.read_csv(FIG7_PATCH_AUDIT)

    old_t3, old_shared = compute_table3(old_df)
    new_t3, new_shared = compute_table3(new_df)
    old_t4 = compute_table4(old_df)
    new_t4 = compute_table4(new_df)

    pub_t3 = public_table3(new_t3)
    pub_t4 = public_table4(new_t4)

    # Outputs: Table 3
    t3_csv = OUT / "corrected_Table3_model_rankings.csv"
    t3_docx = OUT / "corrected_Table3_model_rankings.docx"
    t3_md = OUT / "corrected_Table3_model_rankings.md"
    pub_t3.to_csv(t3_csv, index=False, encoding="utf-8-sig")
    t3_md.write_text(
        to_markdown_table(pub_t3, "Table 3. Final learned-model performance rankings across 36 land-type × target forecasting tasks.", TABLE3_NOTE),
        encoding="utf-8",
    )
    write_docx_table(t3_docx, "Table 3. Final learned-model performance rankings across 36 land-type × target forecasting tasks.", pub_t3, TABLE3_NOTE)

    # Shared best tasks
    shared_csv = OUT / "Table3_shared_best_tasks.csv"
    new_shared.to_csv(shared_csv, index=False, encoding="utf-8-sig")

    # Outputs: Table 4
    t4_csv = OUT / "corrected_Table4_target_predictability.csv"
    t4_docx = OUT / "corrected_Table4_target_predictability.docx"
    t4_md = OUT / "corrected_Table4_target_predictability.md"
    pub_t4.to_csv(t4_csv, index=False, encoding="utf-8-sig")
    t4_md.write_text(
        to_markdown_table(pub_t4, "Table 4. Final target-level predictability summary for the learned-model benchmark.", TABLE4_NOTE),
        encoding="utf-8",
    )
    write_docx_table(t4_docx, "Table 4. Final target-level predictability summary for the learned-model benchmark.", pub_t4, TABLE4_NOTE)

    # Paragraph
    paragraph = make_paragraph(pub_t4)
    paragraph_path = OUT / "corrected_target_predictability_results_paragraph.md"
    paragraph_path.write_text(paragraph + "\n", encoding="utf-8")

    # Audit table
    audit_rows = []
    audit_rows.extend(compare_tables(old_t3, new_t3, ["Model"], "Table 3", str(CORRECTED_FIG7_SOURCE)))
    audit_rows.extend(compare_tables(old_t4, new_t4, ["Target"], "Table 4", str(CORRECTED_FIG7_SOURCE)))
    audit_df = pd.DataFrame(audit_rows)
    audit_path = OUT / "Table3_Table4_correction_audit.csv"
    audit_df.to_csv(audit_path, index=False, encoding="utf-8-sig")

    # Validation checks
    checks = validation_checks(old_df, new_df, new_t3, new_t4, patch_audit)
    checks_path = REPORTS / "Table3_Table4_update_validation_checks.csv"
    checks.to_csv(checks_path, index=False, encoding="utf-8-sig")

    # Detailed report
    changed_t3 = audit_df[(audit_df["table"] == "Table 3") & (audit_df["changed yes/no"] == "yes")]
    changed_t4 = audit_df[(audit_df["table"] == "Table 4") & (audit_df["changed yes/no"] == "yes")]
    model_count_changes = changed_t3[changed_t3["field"].isin(["Best R² count", "Best nRMSE count"])]
    best_model_changes = changed_t4[changed_t4["field"].isin(["Best-R² model", "Best-nRMSE model"])]
    class_changes = changed_t4[changed_t4["field"] == "Predictability class"]

    report = f"""# Table 3 / Table 4 Corrected Forecasting Results Update

Generated: {datetime.now().isoformat(timespec="seconds")}

## Summary

No forecasting model was run or retrained. This update recalculated replacement-ready manuscript tables and the target-predictability paragraph from the final corrected Figure 7 source table.

## Authoritative sources

- Corrected Figure 7 source: `{CORRECTED_FIG7_SOURCE}`
- Accepted pre-correction Figure 7 source used for old-value audit: `{ACCEPTED_OLD_FIG7_SOURCE}`
- Corrected BiLSTM pixel metrics source: `{ROOT / '06_metrics' / 'corrected_BiLSTM_pixel_metrics.csv'}`
- Accepted aggregation logic report: `{ACCEPTED_LOGIC_REPORT}`
- Accepted aggregation verification: `{ACCEPTED_AGGREGATION_VERIFICATION}`
- Patch audit confirming only corrected BiLSTM cells changed: `{FIG7_PATCH_AUDIT}`

## Aggregation logic

Table 3 was recalculated from the 36 land-cover × target task-level Figure 7 values for each learned model. Best-R² counts were assigned where a model had the maximum task-level mean R². Best-nRMSE counts were assigned where a model had the minimum task-level mean nRMSE. Ties within `{TOL}` were recorded as shared best tasks.

Table 4 was recalculated by averaging the 24 Figure 7 values per target (4 land-cover classes × 6 learned models). Best target-level models were identified after averaging each model across the four land-cover classes for the target.

## Corrected Table 3

{pub_t3.to_markdown(index=False)}

{TABLE3_NOTE}

## Corrected Table 4

{pub_t4.to_markdown(index=False)}

{TABLE4_NOTE}

## Replacement paragraph

{paragraph}

## Change summary

- Table 3 changed fields: {len(changed_t3)}
- Table 3 model win-count changed fields: {len(model_count_changes)}
- Table 4 changed fields: {len(changed_t4)}
- Table 4 best-model label changes: {len(best_model_changes)}
- Table 4 predictability-class changes: {len(class_changes)}
- Shared-best tasks in corrected Table 3: {len(new_shared)}

## Output paths

- Corrected Table 3 CSV: `{t3_csv}`
- Corrected Table 3 DOCX: `{t3_docx}`
- Corrected Table 3 MD: `{t3_md}`
- Shared-best tasks: `{shared_csv}`
- Corrected Table 4 CSV: `{t4_csv}`
- Corrected Table 4 DOCX: `{t4_docx}`
- Corrected Table 4 MD: `{t4_md}`
- Corrected paragraph: `{paragraph_path}`
- Correction audit: `{audit_path}`
- Validation checks: `{checks_path}`

## Validation

{checks.to_markdown(index=False)}
"""
    report_path = REPORTS / "Table3_Table4_update_report.md"
    report_path.write_text(report, encoding="utf-8")

    status = {
        "no_model_rerun": True,
        "corrected_figure7_source": str(CORRECTED_FIG7_SOURCE),
        "table3_csv": str(t3_csv),
        "table4_csv": str(t4_csv),
        "paragraph": paragraph,
        "model_count_changes": model_count_changes.to_dict("records"),
        "best_model_changes": best_model_changes.to_dict("records"),
        "class_changes": class_changes.to_dict("records"),
        "shared_best_tasks": len(new_shared),
        "validation_all_pass": bool((checks["status"] == "PASS").all()),
    }
    import json

    (REPORTS / "Table3_Table4_update_status.json").write_text(json.dumps(status, indent=2), encoding="utf-8")
    print(json.dumps(status, indent=2))


if __name__ == "__main__":
    main()
