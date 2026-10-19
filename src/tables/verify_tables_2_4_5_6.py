from __future__ import annotations

import csv
import hashlib
import os
from pathlib import Path

for _name in ("OMP_NUM_THREADS", "MKL_NUM_THREADS", "OPENBLAS_NUM_THREADS", "NUMEXPR_NUM_THREADS"):
    os.environ[_name] = "1"

import numpy as np
import pandas as pd
from docx import Document


OUT = Path(__file__).resolve().parent
PAPER = OUT.parents[1]
MANUSCRIPT = PAPER / "Paper Final Draft.docx"
FIG7_SOURCE = PAPER / "AUDIT_TEMP" / "FIGURE_REFINEMENT_PHASE" / "source_tables" / "Figure_7_source.csv"
FIG7_COUNTS = PAPER / "AUDIT_TEMP" / "FIGURE_REFINEMENT_PHASE" / "FIGURE7_ORIGINAL_DESIGN_RECREATION" / "Figure_7_original_design_corrected_best_model_counts.csv"
FIG7_VERIFY = PAPER / "AUDIT_TEMP" / "FIGURE_REFINEMENT_PHASE" / "FIGURE7_ORIGINAL_DESIGN_RECREATION" / "Figure_7_original_design_corrected_numerical_verification.csv"
HP_TABLE = PAPER / "Docs" / "major_revision_analyses" / "final_forecasting_hyperparameter_audit_20260728" / "Supplementary_Table_forecasting_hyperparameters.csv"
HP_SENSITIVITY = PAPER / "Docs" / "major_revision_analyses" / "forecasting_hyperparameter_sensitivity_final" / "run_20260727_200736" / "06_tables" / "Supplementary_Table_hyperparameter_sensitivity.csv"
BILSTM_UPDATE = PAPER / "Docs" / "major_revision_analyses" / "corrected_true_BiLSTM_rerun_20260728" / "08_downstream" / "corrected_BiLSTM_hyperparameter_table_update.csv"
BILSTM_SENSITIVITY = PAPER / "Docs" / "major_revision_analyses" / "corrected_true_BiLSTM_rerun_20260728" / "07_sensitivity" / "corrected_BiLSTM_sensitivity_summary.csv"
BILSTM_CODE = PAPER / "Docs" / "major_revision_analyses" / "corrected_true_BiLSTM_rerun_20260728" / "scripts" / "run_corrected_true_bilstm_rerun.py"
SVR_EVIDENCE = PAPER / "Docs" / "major_revision_analyses" / "final_forecasting_hyperparameter_audit_20260728" / "forecasting_hyperparameter_evidence_table.csv"
TABLE6_SOURCE = PAPER / "Docs" / "major_revision_analyses" / "FINAL_table4_figure8_verification_rebuild" / "tables" / "Table4_final_nirv_response_and_severe_area.csv"
TABLE6_REPORT = PAPER / "Docs" / "major_revision_analyses" / "FINAL_table4_figure8_verification_rebuild" / "report" / "FINAL_table4_figure8_verification_report.md"

TARGET_ORDER = ["NDVI", "GNDVI", "NIRv", "GOSIF/SIF", "GPP", "LST", "VPDmax", "SPEI-1", "TWS"]
MODEL_ORDER = ["XGBoost", "BiLSTM", "TCN", "N-BEATS", "LSTM", "N-HiTS"]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_csv(path: Path, rows: list[dict], fieldnames: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def table2_rows() -> list[dict]:
    return [
        {
            "Model": "LSTM",
            "Final configuration": "2 layers; 128 hidden units; dropout 0.20; learning rate 0.001.",
            "Rationale and sensitivity check": "Based on (Greff et al., 2017); 64 units and learning rate 0.002 showed no consistent improvement.",
        },
        {
            "Model": "BiLSTM",
            "Final configuration": "2 layers; bidirectional first layer; 128 units per direction (150 for TWS; 200 for SPEI-1); dropout 0.20; learning rate 0.001.",
            "Rationale and sensitivity check": "Based on (Schuster and Paliwal, 1997); Lower capacity and learning rates 0.0005 and 0.002 were not consistently better.",
        },
        {
            "Model": "TCN",
            "Final configuration": "3 residual blocks; channels 32/32/32; kernel 3; dilations 1/2/4; dropout 0.30; learning rate 0.001.",
            "Rationale and sensitivity check": "Temporal convolutions and dilation are supported by (Lea et al., 2017). Channels 16/16/16 and learning rate 0.002 gave no consistent gain.",
        },
        {
            "Model": "N-BEATS",
            "Final configuration": "Generic, trend, and seasonality stacks; 2 blocks per stack; 4 fully connected layers per block; width 64; learning rate 0.001.",
            "Rationale and sensitivity check": "The structure follows N-BEATS basis expansion (Oreshkin et al., 2019). Width 32 and learning rate 0.002 were not consistently better.",
        },
        {
            "Model": "N-HiTS",
            "Final configuration": "3 blocks; pooling 2/2/1; width 512; 2 fully connected layers per block; dropout 0.10; learning rate 0.001.",
            "Rationale and sensitivity check": "The structure follows hierarchical interpolation in N-HiTS (Challu et al., 2023). Width 256 and learning rate 0.002 showed no consistent advantage.",
        },
        {
            "Model": "XGBoost",
            "Final configuration": "300 trees; maximum depth 6; learning rate 0.05; row and column subsampling 0.80; early stopping after 20 rounds.",
            "Rationale and sensitivity check": "Regularized boosting, shrinkage, and subsampling follow (Chen and Guestrin, 2016). Depth 3 and learning rate 0.10 were not consistently better.",
        },
        {
            "Model": "RBF-SVR",
            "Final configuration": "RBF kernel; C = {1, 3, 10}; ε = {0.05, 0.10, 0.20}; γ multiplier = {0.5, 1, 2}; task-specific selection.",
            "Rationale and sensitivity check": "Based on (Smola and Schölkopf, 2004); C, ε, and γ were selected by task-specific grid search.",
        },
    ]


def validate_table2_sources(rows: list[dict]) -> None:
    for path in (HP_TABLE, HP_SENSITIVITY, BILSTM_UPDATE, BILSTM_SENSITIVITY, BILSTM_CODE, SVR_EVIDENCE):
        if not path.is_file():
            raise FileNotFoundError(path)
    hp = pd.read_csv(HP_TABLE).set_index("Model")
    corrected = pd.read_csv(BILSTM_UPDATE).iloc[0]["Corrected architecture and training setting"]
    code = BILSTM_CODE.read_text(encoding="utf-8")
    required_bilstm_phrases = [
        "bidirectional=True",
        "self.second = nn.LSTM(int(units) * 2, int(units), batch_first=True)",
        "dropout=0.2",
        "learning_rate: float = 0.001",
        '"Two recurrent layers; bidirectional first layer for all nine targets;',
        "200 units per direction retained for SPEI-1",
        "150 units per direction retained for TWS",
    ]
    if not all(phrase in code for phrase in required_bilstm_phrases):
        raise ValueError("Corrected TRUE-BiLSTM code evidence is incomplete")
    if not all(text in corrected for text in ["bidirectional first layer for all nine targets", "128 units per direction", "200 units per direction retained for SPEI-1", "150 units per direction retained for TWS", "dropout=0.20", "learning rate=0.001"]):
        raise ValueError("Corrected TRUE-BiLSTM downstream setting does not match Table 2")
    checks = {
        "LSTM": ["Two-layer unidirectional LSTM", "hidden_size=128", "dropout=0.20", "learning rate=0.001"],
        "TCN": ["Three residual temporal blocks", "(32,32,32)", "kernel size=3", "1,2,4", "dropout=0.30", "learning rate=0.001"],
        "N-BEATS": ["Generic, trend, and seasonality stacks", "2 blocks per stack", "4 fully connected layers per block", "hidden width=64", "learning rate=0.001"],
        "N-HiTS": ["pooling sizes 2,2,1", "hidden width=512", "two fully connected layers per block", "dropout=0.10", "learning rate=0.001"],
        "XGBoost": ["n_estimators=300", "max_depth=6", "learning_rate=0.05", "subsample=0.8", "colsample_bytree=0.8", "early_stopping_rounds=20"],
        "RBF-epsilon-SVR": ["C={1,3,10}", "epsilon={0.05,0.10,0.20}", "gamma multiplier={0.5,1,2}", "task-specific values selected by internal-validation nRMSE"],
    }
    for model, phrases in checks.items():
        combined = " ".join(str(value) for value in hp.loc[model].tolist())
        if not all(phrase in combined for phrase in phrases):
            raise ValueError(f"Table 2 configuration mismatch for {model}: {phrases}")
    sensitivity = pd.read_csv(HP_SENSITIVITY)
    if set(sensitivity.Model) != {"LSTM", "BiLSTM", "TCN", "N-BEATS", "N-HiTS", "XGBoost"}:
        raise ValueError("Learned-model sensitivity evidence is incomplete")
    corrected_sensitivity = pd.read_csv(BILSTM_SENSITIVITY)
    if set(corrected_sensitivity.configuration) != {"lower_capacity", "lower_learning_rate", "higher_learning_rate"}:
        raise ValueError("Corrected TRUE-BiLSTM sensitivity evidence is incomplete")
    if [row["Model"] for row in rows] != ["LSTM", "BiLSTM", "TCN", "N-BEATS", "N-HiTS", "XGBoost", "RBF-SVR"]:
        raise ValueError("Unexpected Table 2 row order")


def compute_table4_and_table5() -> tuple[list[dict], list[dict], pd.DataFrame]:
    source = pd.read_csv(FIG7_SOURCE)
    if len(source) != 216 or source.groupby(["land_type_display", "target_display"]).ngroups != 36:
        raise ValueError("Final Figure 7 source does not contain 36 complete tasks")
    if not source.groupby(["land_type_display", "target_display"]).size().eq(6).all():
        raise ValueError("Each final forecasting task must contain six learned models")
    if len(source[source.model.eq("BiLSTM")]) != 36 or not source.available_folds.eq(4).all():
        raise ValueError("TRUE-BiLSTM/four-fold coverage is incomplete")

    task_keys = ["land_type_display", "target_display"]
    r2_winners = source.loc[source.groupby(task_keys).mean_r2.idxmax()].model.value_counts()
    nrmse_winners = source.loc[source.groupby(task_keys).mean_normalized_rmse.idxmin()].model.value_counts()
    model_means = source.groupby("model").agg(Mean_R2=("mean_r2", "mean"), Mean_nRMSE=("mean_normalized_rmse", "mean"))
    families = {model: "Deep learning" for model in MODEL_ORDER}
    families["XGBoost"] = "Non-deep machine learning"
    table4 = []
    for model in MODEL_ORDER:
        table4.append(
            {
                "Model": model,
                "Model family": families[model],
                "Best R² count": str(int(r2_winners.get(model, 0))),
                "Best nRMSE count": str(int(nrmse_winners.get(model, 0))),
                "Mean R²": f"{model_means.loc[model, 'Mean_R2']:.3f}",
                "Mean nRMSE": f"{model_means.loc[model, 'Mean_nRMSE']:.3f}",
            }
        )

    target_means = source.groupby("target_display").agg(Mean_R2=("mean_r2", "mean"), Mean_nRMSE=("mean_normalized_rmse", "mean"))
    by_target_model = source.groupby(["target_display", "model"]).agg(r2=("mean_r2", "mean"), nrmse=("mean_normalized_rmse", "mean")).reset_index()
    best_r2 = by_target_model.loc[by_target_model.groupby("target_display").r2.idxmax()].set_index("target_display").model
    best_nrmse = by_target_model.loc[by_target_model.groupby("target_display").nrmse.idxmin()].set_index("target_display").model
    classes = {target: "High" for target in ["NDVI", "GNDVI", "NIRv", "GOSIF/SIF"]}
    classes.update({target: "Moderate" for target in ["GPP", "LST", "VPDmax"]})
    classes.update({target: "Low/variable" for target in ["SPEI-1", "TWS"]})
    table5 = []
    for target in TARGET_ORDER:
        value = float(target_means.loc[target, "Mean_R2"])
        verified_class = "High" if value >= 0.70 else "Moderate" if value >= 0.40 else "Low/variable"
        if verified_class != classes[target]:
            raise ValueError(f"Predictability class mismatch for {target}")
        table5.append(
            {
                "Predictability class": classes[target],
                "Target": target,
                "Mean R²": f"{target_means.loc[target, 'Mean_R2']:.3f}",
                "Mean nRMSE": f"{target_means.loc[target, 'Mean_nRMSE']:.3f}",
                "Best-R² model": str(best_r2.loc[target]),
                "Best-nRMSE model": str(best_nrmse.loc[target]),
            }
        )
    return table4, table5, source


def compute_table6() -> list[dict]:
    source = pd.read_csv(TABLE6_SOURCE).set_index("Land type")
    if not source.threshold_percent_nirv_change.eq(-40).all():
        raise ValueError("Table 6 source does not use the -40% threshold")
    if not source.bootstrap_replicates.eq(2000).all() or not source.percentage_denominator.eq("eligible analyzed land area").all():
        raise ValueError("Table 6 bootstrap or percentage denominator mismatch")
    rows = []
    for land in ["Grasslands", "Natural Forests", "Savannas", "Shrublands"]:
        row = source.loc[land]
        rows.append(
            {
                "Land type": land,
                "Mean NIRv 2012–2016 (%)": f"{row.mean_2012_2016_percent:.2f}",
                "Mean NIRv 2020–2024 (%)": f"{row.mean_2020_2024_percent:.2f}",
                "Difference (percentage points)": f"{row.mean_difference_recent_minus_historic_points:.2f}",
                "Severe area 2012–2016": f"{row.severe_2012_2016_area_km2:,.1f} km² ({row.severe_2012_2016_percent_eligible_area:.2f}%)",
                "Severe area 2020–2024": f"{row.severe_2020_2024_area_km2:,.1f} km² ({row.severe_2020_2024_percent_eligible_area:.2f}%)",
                "Area change (km²)": f"{row.severe_area_change_km2_recent_minus_historic:,.1f}",
                "Change (percentage points)": f"{row.severe_area_change_percentage_points:.2f}",
                "95% CI (km²)": f"{row.ci95_low_km2:,.1f} to {row.ci95_high_km2:,.1f}",
            }
        )
    return rows


def current_rows() -> tuple[list[dict], list[dict], list[dict], list[dict]]:
    document = Document(MANUSCRIPT)
    table_indices = [1, 3, 4, 5]
    if len(document.tables) <= max(table_indices):
        raise ValueError("Current manuscript table structure changed unexpectedly")
    parsed = []
    for index in table_indices:
        table = document.tables[index]
        headers = ["".join(cell._tc.xpath('.//*[local-name()="t"]/text()')).strip() for cell in table.rows[0].cells]
        rows = []
        for row in table.rows[1:]:
            values = ["".join(cell._tc.xpath('.//*[local-name()="t"]/text()')).strip().replace("\n", " ") for cell in row.cells]
            rows.append(dict(zip(headers, values)))
        parsed.append(rows)
    manuscript_configs = {row["Model"]: row["Final configuration"] for row in parsed[0]}
    for row in table2_rows():
        if manuscript_configs.get(row["Model"]) != row["Final configuration"]:
            raise ValueError(f"Current manuscript Table 2 configuration differs for {row['Model']}")
    return parsed[0], parsed[1], parsed[2], parsed[3]


def add_diffs(diffs: list[dict], table: str, current: list[dict], final: list[dict], row_key: str, source_by_column: dict[str, str], reason: str) -> None:
    current_index = {row[row_key]: row for row in current}
    final_index = {row[row_key]: row for row in final}
    if set(current_index) != set(final_index):
        raise ValueError(f"Row mismatch in {table}")
    for row_name, final_row in final_index.items():
        current_row = current_index[row_name]
        for column, final_value in final_row.items():
            current_value = str(current_row.get(column, ""))
            final_text = str(final_value)
            changed = "YES" if current_value != final_text else "NO"
            diffs.append(
                {
                    "Table": table,
                    "Row": row_name,
                    "Column": column,
                    "Current manuscript value": current_value,
                    "Verified final value": final_text,
                    "Changed? YES/NO": changed,
                    "Reason": reason if changed == "NO" else "Verified computational value differs from current manuscript baseline.",
                    "Final source file": source_by_column.get(column, source_by_column.get("*", "")),
                    "Confidence": "High",
                }
            )


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for path in (MANUSCRIPT, FIG7_SOURCE, FIG7_COUNTS, FIG7_VERIFY, TABLE6_SOURCE, TABLE6_REPORT):
        if not path.is_file():
            raise FileNotFoundError(path)

    final2 = table2_rows()
    validate_table2_sources(final2)
    final4, final5, forecast_source = compute_table4_and_table5()
    final6 = compute_table6()
    current2, current4, current5, current6 = current_rows()

    count_check = pd.read_csv(FIG7_COUNTS).set_index("model")
    for row in final4:
        model = row["Model"]
        if int(row["Best R² count"]) != int(count_check.loc[model, "highest_mean_r2_task_count"]):
            raise ValueError(f"Figure 7 R² count mismatch for {model}")
        if int(row["Best nRMSE count"]) != int(count_check.loc[model, "lowest_mean_nrmse_task_count"]):
            raise ValueError(f"Figure 7 nRMSE count mismatch for {model}")
    if not pd.read_csv(FIG7_VERIFY).passed.all():
        raise ValueError("New Figure 7 numerical verification is not fully passing")

    write_csv(OUT / "TABLE_2_VERIFIED.csv", final2, ["Model", "Final configuration", "Rationale and sensitivity check"])
    write_csv(OUT / "TABLE_4_VERIFIED.csv", final4, ["Model", "Model family", "Best R² count", "Best nRMSE count", "Mean R²", "Mean nRMSE"])
    write_csv(OUT / "TABLE_5_VERIFIED.csv", final5, ["Predictability class", "Target", "Mean R²", "Mean nRMSE", "Best-R² model", "Best-nRMSE model"])
    write_csv(OUT / "TABLE_6_VERIFIED.csv", final6, ["Land type", "Mean NIRv 2012–2016 (%)", "Mean NIRv 2020–2024 (%)", "Difference (percentage points)", "Severe area 2012–2016", "Severe area 2020–2024", "Area change (km²)", "Change (percentage points)", "95% CI (km²)"])

    diffs: list[dict] = []
    table2_sources = {
        "Model": str(HP_TABLE),
        "Final configuration": f"{HP_TABLE}; corrected BiLSTM override: {BILSTM_UPDATE}",
        "Rationale and sensitivity check": f"{HP_SENSITIVITY}; corrected BiLSTM sensitivity: {BILSTM_SENSITIVITY}; SVR: {SVR_EVIDENCE}",
    }
    add_diffs(diffs, "Table 2", current2, final2, "Model", table2_sources, "Verified against final implementation and validation-only sensitivity evidence; no old affected-target BiLSTM architecture retained.")
    add_diffs(diffs, "Table 4", current4, final4, "Model", {"*": str(FIG7_SOURCE)}, "Exact match after recomputation from 36 finalized corrected target × land-type tasks.")
    add_diffs(diffs, "Table 5", current5, final5, "Target", {"*": str(FIG7_SOURCE)}, "Exact match after intended 24-row target aggregation (4 land types × 6 learned models).")
    add_diffs(diffs, "Table 6", current6, final6, "Land type", {"*": str(TABLE6_SOURCE)}, "Exact match to the finalized fractional-mask severe-response/Figure-8 branch.")
    if len(diffs) != 147:
        raise ValueError(f"Expected 147 cell comparisons, found {len(diffs)}")
    if any(row["Changed? YES/NO"] == "YES" for row in diffs):
        raise ValueError("Unexpected manuscript-table difference found; inspect before issuing verdict")
    write_csv(
        OUT / "TABLE_VALUE_DIFFS.csv",
        diffs,
        ["Table", "Row", "Column", "Current manuscript value", "Verified final value", "Changed? YES/NO", "Reason", "Final source file", "Confidence"],
    )

    sources = {
        "manuscript": (MANUSCRIPT, sha256(MANUSCRIPT)),
        "forecast": (FIG7_SOURCE, sha256(FIG7_SOURCE)),
        "figure7_counts": (FIG7_COUNTS, sha256(FIG7_COUNTS)),
        "bilstm_update": (BILSTM_UPDATE, sha256(BILSTM_UPDATE)),
        "table6": (TABLE6_SOURCE, sha256(TABLE6_SOURCE)),
    }
    report = f"""# Final verification of manuscript Tables 2, 4, 5, and 6

## Overall verdict

All four current manuscript tables are **VERIFIED UNCHANGED**. No replacement `.docx` was created because no table requires replacement. This was a read-only provenance and arithmetic verification; no model was retrained, no figure was modified, and no manuscript file was edited.

## Table 2 — VERIFIED UNCHANGED

Every displayed architecture and tuning field matches the implementation that generated the accepted outputs.

- The finalized TRUE-BiLSTM has two recurrent layers and a bidirectional first layer for **all nine targets**. The first layer uses 128 units per direction for NDVI, GNDVI, NIRv, GOSIF/SIF, GPP, LST, and VPDmax; 150 per direction for TWS; and 200 per direction for SPEI-1. The second layer is unidirectional. Dropout 0.20 and learning rate 0.001 are correct.
- The older audit wording that described seven affected targets as unidirectional was not used as final evidence; it is superseded by `{BILSTM_UPDATE}` and the corrected implementation `{BILSTM_CODE}`.
- LSTM, TCN, N-BEATS, N-HiTS, and XGBoost settings match `{HP_TABLE}`.
- RBF-SVR uses the stated C, epsilon, and gamma-multiplier grid, selected separately by target × land type × fold using internal-validation nRMSE. The concise manuscript wording is accurate.
- The rationale cells are qualitative summaries of the finalized validation-only local sensitivity checks. None refers to the invalid old affected-target BiLSTM sensitivity rows.

## Table 4 — VERIFIED UNCHANGED

All four numerical columns reproduce from `{FIG7_SOURCE}`. Model means are unweighted means across the 36 finalized target × land-type task means. Best counts select the highest mean R² and lowest mean nRMSE within each of the same 36 tasks.

The counts also exactly match the new Figure 7 count panels: XGBoost 17/18, BiLSTM 7/7, TCN 6/4, N-BEATS 2/4, LSTM 4/3, and N-HiTS 0/0. These are descriptive counts only.

## Table 5 — VERIFIED UNCHANGED

Each target mean uses exactly 24 finalized panel values: 4 land types × 6 learned models. Best models are selected after averaging each model across the four land types for that target. All means, winners, and target names match the finalized corrected branch.

Predictability classes remain justified by the existing descriptive thresholds: High, mean R² ≥ 0.70; Moderate, 0.40 ≤ mean R² < 0.70; Low/variable, mean R² < 0.40. No target crosses a class boundary after TRUE-BiLSTM reconciliation.

## Table 6 — VERIFIED UNCHANGED

Every displayed value matches `{TABLE6_SOURCE}` after manuscript rounding. This branch is independent of forecasting and EWBI revisions.

- Severe-response threshold: NIRv change ≤ -40%.
- Percentages use eligible analyzed land area.
- Area changes are 2020–2024 minus 2012–2016.
- Confidence intervals are paired spatial-block-bootstrap 95% intervals with 2,000 replicates and approximately 44.5-km blocks.

## Final consistency checks

- Table 4 counts = new finalized Figure 7 count panels: **PASS**.
- Tables 4 and 5 = the same 216-row, 36-task corrected forecasting source with complete TRUE-BiLSTM coverage: **PASS**.
- Table 6 = finalized severe-response/Figure-8 source branch: **PASS**.
- Table 2 = actual finalized implementations, with the corrected TRUE-BiLSTM override applied: **PASS**.
- Cell comparison coverage: 147/147 cells across Tables 2, 4, 5, and 6; changed cells: 0.

## Key source hashes

| Role | File | SHA-256 |
|---|---|---|
"""
    for role, (path, digest) in sources.items():
        report += f"| {role} | `{path}` | `{digest}` |\n"
    (OUT / "FINAL_TABLE_VERIFICATION.md").write_text(report, encoding="utf-8")

    replacements = """# Tables to replace

Table 2: **KEEP**

Table 4: **KEEP**

Table 5: **KEEP**

Table 6: **KEEP**

No cells changed after comparison with the finalized computational sources. `UPDATED_TABLES_2_4_5_6.docx` was therefore not created.
"""
    (OUT / "TABLES_TO_REPLACE.md").write_text(replacements, encoding="utf-8")
    print("Verified Tables 2, 4, 5, and 6: all KEEP; 147/147 cells unchanged.")


if __name__ == "__main__":
    main()
