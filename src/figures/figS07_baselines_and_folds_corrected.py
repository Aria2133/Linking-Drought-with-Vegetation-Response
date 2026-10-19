from __future__ import annotations

import hashlib
import json
from datetime import datetime
from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.gridspec import GridSpecFromSubplotSpec
from matplotlib.lines import Line2D
from matplotlib.patches import Patch, Rectangle
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"<PROJECT_ROOT>\Docs\major_revision_analyses\corrected_true_BiLSTM_rerun_20260728")
DOWNSTREAM = ROOT / "08_downstream"
REPORTS = ROOT / "09_reports"

FIG7_ACCEPTED_ROOT = Path(r"<PROJECT_ROOT>\Docs\major_revision_analyses\ALTERNATIVE_Figure7_polar_redesign_no_detrend_update")
FIG7_ACCEPTED_TABLE = FIG7_ACCEPTED_ROOT / "Source tables used for plotting" / "figure7_alternative_plotting_table.csv"
FIG7_ACCEPTED_REPORT = FIG7_ACCEPTED_ROOT / "Report" / "Figure7_alternative_rebuild_report.md"
FIG7_ACCEPTED_V2_REPORT = FIG7_ACCEPTED_ROOT / "Report" / "Figure7_alternative_large_readable_v2_update_report.md"
FIG7_ACCEPTED_V2_PDF = FIG7_ACCEPTED_ROOT / "Figures" / "Figure_7_alternative_polar_large_readable_v2.pdf"
FIG7_ACCEPTED_PNG = FIG7_ACCEPTED_ROOT / "Figures" / "Figure_7_alternative_polar_large_readable.png"

CORRECTED_PIXEL = ROOT / "06_metrics" / "corrected_BiLSTM_pixel_metrics.csv"
CORRECTED_MONTHLY = ROOT / "06_metrics" / "corrected_BiLSTM_monthly_metrics.csv"

S6_ACCEPTED_ROOT = Path(r"<PROJECT_ROOT>\Linking-Drought-with-Vegetation-Response\outputs\supplementary_figure_S6_update_with_SVR")
S6_ACCEPTED_RENDERER = S6_ACCEPTED_ROOT / "scripts" / "build_figure_s6_update_with_svr.py"
S6_ACCEPTED_FIG = Path(r"<PROJECT_ROOT>\Paper Revised Version\Figure_S6_updated_with_SVR.png")
S6_ACCEPTED_MONTHLY = S6_ACCEPTED_ROOT / "source_tables" / "Figure_S6_updated_combined_monthly_metrics_with_SVR.csv"
S6_ACCEPTED_PIXEL = S6_ACCEPTED_ROOT / "source_tables" / "Figure_S6_updated_combined_pixel_metrics_learned_source.csv"
S6_ACCEPTED_PANEL_A = S6_ACCEPTED_ROOT / "source_tables" / "Figure_S6_updated_monthly_aggregate_benchmark_summary.csv"
S6_ACCEPTED_PANEL_B = S6_ACCEPTED_ROOT / "source_tables" / "Figure_S6_updated_fold_performance_summary.csv"
S6_ACCEPTED_PANEL_C = S6_ACCEPTED_ROOT / "source_tables" / "Figure_S6_rolling_origin_timeline_data.csv"

SENS_ROOT = Path(r"<PROJECT_ROOT>\Docs\major_revision_analyses\forecasting_hyperparameter_sensitivity_final\run_20260727_200736")
SENS_SUMMARY = SENS_ROOT / "06_tables" / "compact_hyperparameter_sensitivity_summary.csv"
SENS_TABLE = SENS_ROOT / "06_tables" / "Supplementary_Table_hyperparameter_sensitivity.csv"
SENS_SVR_SUMMARY = SENS_ROOT / "06_tables" / "SVR_hyperparameter_sensitivity_existing_grid_summary.csv"
SENS_SVR_SELECTED = SENS_ROOT / "05_svr" / "svr_gndvi_spei_no_detrend_selected_configurations.csv"
SENS_REPORT = SENS_ROOT / "09_reports" / "hyperparameter_sensitivity_report.md"
CORR_BILSTM_SENS_SUMMARY = ROOT / "07_sensitivity" / "corrected_BiLSTM_sensitivity_summary.csv"
CORR_BILSTM_SENS_RAW = ROOT / "07_sensitivity" / "corrected_BiLSTM_sensitivity_raw_results.csv"
INVALID_OLD_BILSTM = ROOT / "07_sensitivity" / "invalid_provisional_old_BiLSTM_sensitivity_rows.csv"
FINAL_HPARAM = Path(r"<PROJECT_ROOT>\Docs\major_revision_analyses\final_forecasting_hyperparameter_audit_20260728\Supplementary_Table_forecasting_hyperparameters.csv")
CORR_BILSTM_HPARAM = ROOT / "08_downstream" / "corrected_BiLSTM_hyperparameter_table_update.csv"

TARGET_ORDER = ["GNDVI", "GOSIF/SIF", "GPP", "LST", "NDVI", "NIRv", "SPEI-1", "TWS", "VPDmax"]
LAND_ORDER = ["Grasslands", "Natural Forests", "Savannas", "Shrublands"]
LAND_INTERNAL = {
    "Grasslands": "GRASSLANDS",
    "Natural Forests": "NATURAL_FORESTS",
    "Savannas": "SAVANNAS",
    "Shrublands": "SHRUBLANDS",
}
MODEL_ORDER = ["LSTM", "BiLSTM", "TCN", "N-BEATS", "N-HiTS", "XGBoost"]
S6_MODEL_ORDER = [
    "Persistence",
    "Seasonal persistence",
    "Monthly climatology",
    "RBF-epsilon-SVR",
    "LSTM",
    "BiLSTM",
    "TCN",
    "N-BEATS",
    "N-HiTS",
    "XGBoost",
]
PATCH_INTERNAL_TARGETS = {"NDVI", "GNDVI", "NIRv", "GOSIF", "GPP", "LST", "VPD"}
DISPLAY = {"GOSIF": "GOSIF/SIF", "VPD": "VPDmax", "SPEI": "SPEI-1"}
INTERNAL = {v: k for k, v in DISPLAY.items()}
INTERNAL.update({"NDVI": "NDVI", "GNDVI": "GNDVI", "NIRv": "NIRv", "GPP": "GPP", "LST": "LST", "TWS": "TWS"})

MODEL_COLORS = {
    "LSTM": "#1f77b4",
    "BiLSTM": "#9467bd",
    "TCN": "#2ca02c",
    "N-BEATS": "#ff7f0e",
    "N-HiTS": "#d62728",
    "XGBoost": "#17becf",
}
S6_COLORS = {
    "Persistence": "#6F6F6F",
    "Seasonal persistence": "#A3A3A3",
    "Monthly climatology": "#B8A900",
    "RBF-epsilon-SVR": "#4DAF4A",
    "LSTM": "#2F6DAE",
    "BiLSTM": "#8556A8",
    "TCN": "#3B8C4A",
    "N-BEATS": "#F28E2B",
    "N-HiTS": "#D84A4A",
    "XGBoost": "#00A6A6",
}


def sha256(path: Path) -> str:
    if not path.exists():
        return ""
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def configure_common_style() -> None:
    mpl.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "figure.facecolor": "white",
            "savefig.facecolor": "white",
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
            "axes.facecolor": "white",
        }
    )


def parse_yyyymm(value: object) -> float:
    text = str(value)
    if "-" in text:
        year, month = text.split("-")[:2]
    else:
        year, month = text[:4], text[4:6]
    return int(year) + (int(month) - 1) / 12


def panel_label(ax: plt.Axes, label: str, x: float = -0.10, y: float = 1.07) -> None:
    ax.text(x, y, label, transform=ax.transAxes, fontsize=12.5, fontweight="bold", va="top", ha="left", clip_on=False)


def aggregate_corrected_pixel() -> pd.DataFrame:
    metrics = pd.read_csv(CORRECTED_PIXEL)
    metrics = metrics[(metrics["model"] == "BiLSTM") & metrics["variable"].isin(PATCH_INTERNAL_TARGETS)].copy()
    grouped = (
        metrics.groupby(["variable", "land_type"], as_index=False)
        .agg(
            mean_r2=("r2", "mean"),
            mean_normalized_rmse=("normalized_rmse", "mean"),
            sd_r2=("r2", "std"),
            sd_normalized_rmse=("normalized_rmse", "std"),
            available_rows=("r2", "size"),
            available_folds=("fold", "nunique"),
            available_seeds=("seed", "nunique"),
            selected_configuration=("preprocessing", lambda x: ";".join(sorted(set(map(str, x))))),
        )
    )
    grouped["model"] = "BiLSTM"
    grouped["target_internal"] = grouped["variable"]
    grouped["target_display"] = grouped["variable"].map(lambda x: DISPLAY.get(x, x))
    grouped["land_type_internal"] = grouped["land_type"]
    inv_land = {v: k for k, v in LAND_INTERNAL.items()}
    grouped["land_type_display"] = grouped["land_type"].map(inv_land)
    grouped["data_status"] = "patched_corrected_true_BiLSTM"
    grouped["source_metric_table"] = str(CORRECTED_PIXEL)
    grouped["metric_uncertainty_definition"] = "SD across available seed/fold metric rows from corrected true-BiLSTM source table convention"
    grouped["r2_negative_flag"] = grouped["mean_r2"] < 0
    grouped["r2_plot_radius"] = grouped["mean_r2"].clip(lower=0)
    grouped["nrmse_plot_radius"] = grouped["mean_normalized_rmse"]
    grouped["r2_plot_note"] = np.where(grouped["r2_negative_flag"], "negative R2 shown on zero ring in polar plot; exact value retained in mean_r2", "")
    return grouped


def patch_figure7_source() -> tuple[pd.DataFrame, pd.DataFrame, int]:
    accepted = pd.read_csv(FIG7_ACCEPTED_TABLE)
    corrected = aggregate_corrected_pixel()
    patched = accepted.copy()

    lookup = corrected.set_index(["target_internal", "land_type_internal"]).to_dict("index")
    audit_rows: list[dict[str, object]] = []
    replaced_keys = set()
    for idx, row in patched.iterrows():
        target_internal = str(row["target_internal"])
        land_internal = str(row["land_type_internal"])
        is_patch = row["model"] == "BiLSTM" and target_internal in PATCH_INTERNAL_TARGETS
        source_path = str(row["source_metric_table"])
        if is_patch:
            new = lookup[(target_internal, land_internal)]
            replaced_keys.add((target_internal, land_internal, "BiLSTM"))
            source_path = str(CORRECTED_PIXEL)
            for col in [
                "mean_r2",
                "mean_normalized_rmse",
                "sd_r2",
                "sd_normalized_rmse",
                "available_rows",
                "available_folds",
                "available_seeds",
                "selected_configuration",
                "data_status",
                "source_metric_table",
                "metric_uncertainty_definition",
                "r2_negative_flag",
                "r2_plot_radius",
                "nrmse_plot_radius",
                "r2_plot_note",
            ]:
                patched.at[idx, col] = new[col]
        for metric_col in ["mean_r2", "mean_normalized_rmse"]:
            old_value = float(row[metric_col])
            corrected_value = float(patched.at[idx, metric_col])
            audit_rows.append(
                {
                    "land_cover": row["land_type_display"],
                    "target": row["target_display"],
                    "model": row["model"],
                    "metric": "R2" if metric_col == "mean_r2" else "nRMSE",
                    "old_accepted_value": old_value,
                    "corrected_value": corrected_value,
                    "changed": abs(old_value - corrected_value) > 1e-12,
                    "source_path": source_path,
                }
            )

    if len(replaced_keys) != 28:
        raise RuntimeError(f"Expected 28 Figure 7 BiLSTM rows to patch, found {len(replaced_keys)}")
    patched.to_csv(DOWNSTREAM / "final_corrected_Figure_7_source_table.csv", index=False)
    audit = pd.DataFrame(audit_rows)
    audit.to_csv(DOWNSTREAM / "Figure7_BiLSTM_patch_audit.csv", index=False)
    return accepted, patched, len(replaced_keys)


def render_polar_figure(table: pd.DataFrame, png: Path, pdf: Path) -> None:
    configure_common_style()
    target_angles = np.linspace(0, 2 * np.pi, len(TARGET_ORDER), endpoint=False)
    offset_width = 0.105
    offsets = np.linspace(-offset_width, offset_width, len(MODEL_ORDER))

    fig, axes = plt.subplots(
        nrows=4,
        ncols=2,
        subplot_kw={"projection": "polar"},
        figsize=(22.46, 40.37),
        dpi=500,
    )
    plt.subplots_adjust(left=0.055, right=0.945, top=0.975, bottom=0.055, wspace=0.16, hspace=0.23)

    for row_idx, land in enumerate(LAND_ORDER):
        for col_idx, metric in enumerate(["R2", "nRMSE"]):
            ax = axes[row_idx, col_idx]
            ax.set_theta_offset(np.pi / 2)
            ax.set_theta_direction(-1)
            ax.set_xticks(target_angles)
            ax.set_xticklabels(TARGET_ORDER, fontsize=18, fontweight="bold")
            ax.grid(True, color="#CFCFCF", alpha=0.55, linewidth=1.0)
            ax.spines["polar"].set_color("#333333")
            ax.spines["polar"].set_linewidth(1.6)
            ax.set_rlabel_position(225)
            if metric == "R2":
                ax.set_ylim(0, 1.0)
                ax.set_yticks([0, 0.25, 0.5, 0.75, 1.0])
                ax.set_yticklabels(["0", "0.25", "0.50", "0.75", "1"], fontsize=15)
                title = f"{land} - R$^2$"
                value_col, err_col = "r2_plot_radius", "sd_r2"
            else:
                ax.set_ylim(0, 2.0)
                ax.set_yticks([0, 0.5, 1.0, 1.5, 2.0])
                ax.set_yticklabels(["0", "0.5", "1.0", "1.5", "2"], fontsize=15)
                title = f"{land} - nRMSE"
                value_col, err_col = "nrmse_plot_radius", "sd_normalized_rmse"
            ax.set_title(title, fontsize=24, fontweight="bold", pad=24)

            sub = table[table["land_type_display"] == land]
            for model_idx, model in enumerate(MODEL_ORDER):
                model_sub = sub[sub["model"] == model].set_index("target_display")
                color = MODEL_COLORS[model]
                for target_idx, target in enumerate(TARGET_ORDER):
                    if target not in model_sub.index:
                        continue
                    data = model_sub.loc[target]
                    theta = target_angles[target_idx] + offsets[model_idx]
                    radius = float(data[value_col])
                    err = float(data[err_col]) if pd.notna(data[err_col]) else np.nan
                    if metric == "R2" and bool(data["r2_negative_flag"]):
                        ax.scatter(theta, 0.02, s=85, marker="x", color=color, linewidths=2.4, zorder=5)
                    else:
                        yerr = None
                        if np.isfinite(err) and err > 0:
                            yerr = [[min(err, max(radius, 0))], [err]]
                        ax.errorbar(
                            theta,
                            radius,
                            yerr=yerr,
                            fmt="o",
                            markersize=10.5,
                            markerfacecolor=color,
                            markeredgecolor="white",
                            markeredgewidth=0.9,
                            ecolor=color,
                            elinewidth=2.0,
                            capsize=4.2,
                            alpha=0.95,
                            zorder=4,
                        )

    handles = [
        Line2D([0], [0], marker="o", linestyle="None", markersize=13, markerfacecolor=MODEL_COLORS[m], markeredgecolor="white", label=m)
        for m in MODEL_ORDER
    ]
    legend = fig.legend(
        handles=handles,
        loc="lower center",
        bbox_to_anchor=(0.5, 0.024),
        ncol=6,
        title="Learned models",
        frameon=True,
        fancybox=False,
        edgecolor="#333333",
        facecolor="white",
        fontsize=17,
        title_fontsize=18,
        borderpad=0.9,
        labelspacing=0.9,
        handletextpad=0.7,
        columnspacing=1.8,
    )
    legend.get_title().set_fontweight("bold")
    fig.text(
        0.5,
        0.006,
        "Radius gives the plotted metric value. Error bars are SD across available seed/fold metric rows. x at center = mean R^2 < 0.",
        ha="center",
        fontsize=13,
        color="#333333",
    )
    fig.savefig(png, dpi=500, bbox_inches="tight")
    fig.savefig(pdf, bbox_inches="tight")
    plt.close(fig)


def patch_s6_sources() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, int, int]:
    monthly_old = pd.read_csv(S6_ACCEPTED_MONTHLY)
    pixel_old = pd.read_csv(S6_ACCEPTED_PIXEL)
    monthly_corr = pd.read_csv(CORRECTED_MONTHLY)
    pixel_corr = pd.read_csv(CORRECTED_PIXEL)
    monthly_corr = monthly_corr[(monthly_corr["model"] == "BiLSTM") & monthly_corr["variable"].isin(PATCH_INTERNAL_TARGETS)].copy()
    pixel_corr = pixel_corr[(pixel_corr["model"] == "BiLSTM") & pixel_corr["variable"].isin(PATCH_INTERNAL_TARGETS)].copy()

    key_cols = ["model", "variable", "land_type", "fold", "seed"]
    def patch_rows(old: pd.DataFrame, corr: pd.DataFrame) -> pd.DataFrame:
        keep = ~(
            (old["model"] == "BiLSTM")
            & (old["variable"].isin(PATCH_INTERNAL_TARGETS))
        )
        out = pd.concat([old[keep].copy(), corr.reindex(columns=old.columns).copy()], ignore_index=True, sort=False)
        return out.sort_values([c for c in key_cols if c in out.columns]).reset_index(drop=True)

    monthly_new = patch_rows(monthly_old, monthly_corr)
    pixel_new = patch_rows(pixel_old, pixel_corr)
    monthly_new.to_csv(DOWNSTREAM / "final_corrected_Figure_S6_combined_monthly_metrics_with_SVR.csv", index=False)
    pixel_new.to_csv(DOWNSTREAM / "final_corrected_Figure_S6_combined_pixel_metrics_learned_source.csv", index=False)

    panel_a_old = make_s6_panel_a(monthly_old)
    panel_b_old = make_s6_panel_b(pixel_old)
    panel_a = make_s6_panel_a(monthly_new)
    panel_b = make_s6_panel_b(pixel_new)
    panel_c = pd.read_csv(S6_ACCEPTED_PANEL_C)
    panel_a.to_csv(DOWNSTREAM / "final_corrected_Figure_S6_panel_a_source.csv", index=False)
    panel_b.to_csv(DOWNSTREAM / "final_corrected_Figure_S6_panel_b_source.csv", index=False)
    panel_c.to_csv(DOWNSTREAM / "final_corrected_Figure_S6_panel_c_source.csv", index=False)

    audit_rows: list[dict[str, object]] = []
    for model in S6_MODEL_ORDER:
        old_row = panel_a_old[panel_a_old["model"] == model].iloc[0]
        new_row = panel_a[panel_a["model"] == model].iloc[0]
        for metric in ["mean_r2", "sd_r2", "mean_normalized_rmse", "sd_normalized_rmse"]:
            audit_rows.append(
                {
                    "panel": "a",
                    "model": model,
                    "fold": "",
                    "metric": metric,
                    "old_accepted_value": old_row[metric],
                    "corrected_value": new_row[metric],
                    "changed": abs(float(old_row[metric]) - float(new_row[metric])) > 1e-12,
                    "source_path": str(CORRECTED_MONTHLY) if model == "BiLSTM" else str(S6_ACCEPTED_MONTHLY),
                }
            )
    for fold in sorted(panel_b["fold"].unique()):
        old_row = panel_b_old[panel_b_old["fold"] == fold].iloc[0]
        new_row = panel_b[panel_b["fold"] == fold].iloc[0]
        for metric in ["mean_r2", "sd_r2", "mean_nrmse", "sd_nrmse"]:
            audit_rows.append(
                {
                    "panel": "b",
                    "model": "learned-model aggregate",
                    "fold": int(fold),
                    "metric": metric,
                    "old_accepted_value": old_row[metric],
                    "corrected_value": new_row[metric],
                    "changed": abs(float(old_row[metric]) - float(new_row[metric])) > 1e-12,
                    "source_path": f"{S6_ACCEPTED_PIXEL}; patched BiLSTM source: {CORRECTED_PIXEL}",
                }
            )
    for _, row in panel_c.iterrows():
        audit_rows.append(
            {
                "panel": "c",
                "model": "rolling-origin schematic",
                "fold": int(row["fold"]),
                "metric": "timeline",
                "old_accepted_value": json.dumps(row.to_dict(), sort_keys=True),
                "corrected_value": json.dumps(row.to_dict(), sort_keys=True),
                "changed": False,
                "source_path": str(S6_ACCEPTED_PANEL_C),
            }
        )
    audit = pd.DataFrame(audit_rows)
    audit.to_csv(DOWNSTREAM / "FigureS6_BiLSTM_patch_audit.csv", index=False)
    return monthly_new, pixel_new, panel_a, panel_b, panel_c, len(monthly_corr), len(pixel_corr)


def make_s6_panel_a(monthly: pd.DataFrame) -> pd.DataFrame:
    sub = monthly[monthly["model"].isin(S6_MODEL_ORDER) & monthly["r2"].notna() & monthly["normalized_rmse"].notna()].copy()
    task = sub.groupby(["model", "land_type", "variable", "fold"], as_index=False, observed=True)[["r2", "normalized_rmse"]].mean()
    summary = (
        task.groupby("model", as_index=False, observed=True)
        .agg(
            n_tasks=("r2", "size"),
            mean_r2=("r2", "mean"),
            sd_r2=("r2", "std"),
            mean_normalized_rmse=("normalized_rmse", "mean"),
            sd_normalized_rmse=("normalized_rmse", "std"),
        )
        .set_index("model")
        .reindex(S6_MODEL_ORDER)
        .reset_index()
    )
    return summary


def make_s6_panel_b(pixel: pd.DataFrame) -> pd.DataFrame:
    sub = pixel[pixel["model"].isin(MODEL_ORDER) & pixel["r2"].notna() & pixel["normalized_rmse"].notna()].copy()
    task = sub.groupby(["model", "land_type", "variable", "fold"], as_index=False, observed=True)[["r2", "normalized_rmse"]].mean()
    folds = (
        task.groupby("fold", as_index=False, observed=True)
        .agg(
            mean_r2=("r2", "mean"),
            sd_r2=("r2", "std"),
            mean_nrmse=("normalized_rmse", "mean"),
            sd_nrmse=("normalized_rmse", "std"),
            n_tasks=("r2", "size"),
        )
        .sort_values("fold")
    )
    folds["test_period"] = folds["fold"].map({1: "2017-18", 2: "2019-20", 3: "2021-22", 4: "2023-24"})
    return folds


def render_s6(panel_a: pd.DataFrame, panel_b: pd.DataFrame, panel_c: pd.DataFrame, png: Path, pdf: Path) -> None:
    configure_common_style()
    mpl.rcParams.update(
        {
            "font.size": 10.0,
            "axes.titlesize": 11.5,
            "axes.titleweight": "semibold",
            "axes.labelsize": 10.0,
            "xtick.labelsize": 9.0,
            "ytick.labelsize": 9.0,
            "legend.fontsize": 8.4,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.linewidth": 0.8,
            "grid.color": "#D6D6D6",
            "grid.linewidth": 0.55,
            "grid.alpha": 0.55,
        }
    )
    fig = plt.figure(figsize=(13.0, 7.3), layout="constrained")
    gs = fig.add_gridspec(2, 2, height_ratios=[1.15, 1.0])
    sub_a = GridSpecFromSubplotSpec(1, 2, subplot_spec=gs[0, 0], wspace=0.12)
    y = np.arange(len(panel_a))
    colors_model = [S6_COLORS[m] for m in panel_a["model"]]

    ax_a1 = fig.add_subplot(sub_a[0, 0])
    ax_a1.errorbar(panel_a["mean_r2"], y, xerr=panel_a["sd_r2"], fmt="none", ecolor="#3A3A3A", capsize=2.8, lw=1.0)
    ax_a1.scatter(panel_a["mean_r2"], y, c=colors_model, s=38, zorder=3)
    ax_a1.axvline(0, color="#8B1A1A", lw=0.85)
    ax_a1.set_yticks(y, panel_a["model"])
    ax_a1.invert_yaxis()
    ax_a1.set_xlabel(r"Mean $R^2$")
    ax_a1.set_title(r"Monthly-aggregate $R^2$")
    ax_a1.grid(axis="x")
    panel_label(ax_a1, "(a)", x=-0.28)

    ax_a2 = fig.add_subplot(sub_a[0, 1], sharey=ax_a1)
    ax_a2.errorbar(panel_a["mean_normalized_rmse"], y, xerr=panel_a["sd_normalized_rmse"], fmt="none", ecolor="#3A3A3A", capsize=2.8, lw=1.0)
    ax_a2.scatter(panel_a["mean_normalized_rmse"], y, c=colors_model, s=38, zorder=3)
    ax_a2.set_xlabel("Mean nRMSE")
    ax_a2.set_title("Monthly-aggregate nRMSE")
    ax_a2.tick_params(labelleft=False)
    ax_a2.grid(axis="x")

    sub_b = GridSpecFromSubplotSpec(1, 2, subplot_spec=gs[0, 1], wspace=0.28)
    xb = np.arange(len(panel_b))
    ax_b1 = fig.add_subplot(sub_b[0, 0])
    ax_b1.errorbar(xb, panel_b["mean_r2"], yerr=panel_b["sd_r2"], fmt="o-", ms=5.2, capsize=2.8, lw=1.2, color="#2F6DAE")
    ax_b1.axhline(0, color="#4A4A4A", lw=0.8)
    ax_b1.set_xticks(xb, panel_b["test_period"])
    ax_b1.set_ylabel(r"Mean pixel-level $R^2$")
    ax_b1.set_title(r"Pixel-level $R^2$ by fold")
    ax_b1.grid(axis="y")
    panel_label(ax_b1, "(b)", x=0.03, y=0.98)

    ax_b2 = fig.add_subplot(sub_b[0, 1])
    ax_b2.errorbar(xb, panel_b["mean_nrmse"], yerr=panel_b["sd_nrmse"], fmt="s-", ms=5.2, capsize=2.8, lw=1.2, color="#D84A4A")
    ax_b2.set_xticks(xb, panel_b["test_period"])
    ax_b2.set_ylabel("Mean pixel-level nRMSE")
    ax_b2.set_title("Pixel-level nRMSE by fold")
    ax_b2.grid(axis="y")

    ax_c = fig.add_subplot(gs[1, :])
    ax_c.set_xlim(2000.5, 2025.0)
    ax_c.set_ylim(0.4, 4.6)
    ax_c.set_yticks([1, 2, 3, 4], [f"Fold {i}" for i in [1, 2, 3, 4]])
    for row in panel_c.itertuples(index=False):
        fold = int(row.fold)
        core_end = parse_yyyymm(row.core_model_training_end) + 1 / 12
        val_start = parse_yyyymm(row.internal_validation_start)
        val_end = parse_yyyymm(row.internal_validation_end) + 1 / 12
        test_start = parse_yyyymm(row.test_start)
        test_end = parse_yyyymm(row.test_end) + 1 / 12
        ax_c.add_patch(Rectangle((2001.0, fold - 0.20), core_end - 2001.0, 0.40, facecolor="#8A8A8A", edgecolor="none"))
        ax_c.add_patch(Rectangle((val_start, fold - 0.20), val_end - val_start, 0.40, facecolor="#E9B949", edgecolor="none"))
        ax_c.add_patch(Rectangle((test_start, fold - 0.20), test_end - test_start, 0.40, facecolor="#D95D39", edgecolor="none"))
    ax_c.set_xticks([2001, 2005, 2009, 2013, 2017, 2021, 2025])
    ax_c.set_xlabel("Calendar year")
    ax_c.set_title("Rolling-origin design")
    ax_c.grid(axis="x")
    ax_c.legend(
        handles=[
            Patch(facecolor="#8A8A8A", label="Training"),
            Patch(facecolor="#E9B949", label="Validation"),
            Patch(facecolor="#D95D39", label="Test"),
        ],
        frameon=False,
        ncol=3,
        loc="lower center",
    )
    panel_label(ax_c, "(c)", x=-0.03)
    fig.savefig(png, dpi=600, bbox_inches="tight")
    fig.savefig(pdf, bbox_inches="tight")
    plt.close(fig)


def make_side_by_side(left: Path, right: Path, out: Path, left_label: str, right_label: str, max_height: int = 2200) -> None:
    Image.MAX_IMAGE_PIXELS = None
    img_l = Image.open(left).convert("RGB")
    img_r = Image.open(right).convert("RGB")
    h = min(max_height, img_l.height, img_r.height)
    def resize_to_h(img: Image.Image) -> Image.Image:
        w = int(img.width * h / img.height)
        return img.resize((w, h), Image.Resampling.LANCZOS)
    img_l = resize_to_h(img_l)
    img_r = resize_to_h(img_r)
    pad = 36
    label_h = 70
    canvas = Image.new("RGB", (img_l.width + img_r.width + pad * 3, h + label_h + pad), "white")
    canvas.paste(img_l, (pad, label_h))
    canvas.paste(img_r, (img_l.width + pad * 2, label_h))
    draw = ImageDraw.Draw(canvas)
    try:
        font = ImageFont.truetype("arial.ttf", 28)
    except Exception:
        font = ImageFont.load_default()
    draw.text((pad, 18), left_label, fill="black", font=font)
    draw.text((img_l.width + pad * 2, 18), right_label, fill="black", font=font)
    canvas.save(out, dpi=(200, 200))


def create_docx_from_markdown(md_path: Path, docx_path: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    for line in md_path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line.startswith("|"):
            # Tables are supplied as standalone DOCX; keep the markdown table text compact here.
            doc.add_paragraph(line)
        else:
            doc.add_paragraph(line)
    doc.save(docx_path)


def csv_to_docx_table(csv_path: Path, docx_path: Path, title: str) -> None:
    from docx import Document
    from docx.shared import Pt

    df = pd.read_csv(csv_path)
    doc = Document()
    doc.add_heading(title, level=1)
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Table Grid"
    for i, col in enumerate(df.columns):
        table.rows[0].cells[i].text = str(col)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, col in enumerate(df.columns):
            value = row[col]
            if isinstance(value, float):
                cells[i].text = f"{value:.6g}"
            else:
                cells[i].text = "" if pd.isna(value) else str(value)
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(9)
    doc.save(docx_path)


def build_sensitivity_status_report() -> tuple[Path, Path, Path, Path, Path]:
    compact = pd.read_csv(SENS_SUMMARY)
    compact_table = pd.read_csv(SENS_TABLE)
    svr = pd.read_csv(SENS_SVR_SUMMARY)
    corr = pd.read_csv(CORR_BILSTM_SENS_SUMMARY)
    invalid_count = len(pd.read_csv(INVALID_OLD_BILSTM)) if INVALID_OLD_BILSTM.exists() else 0
    hparam = pd.read_csv(FINAL_HPARAM)
    bilstm_hparam = pd.read_csv(CORR_BILSTM_HPARAM)

    rows: list[dict[str, object]] = []
    for _, row in compact_table.iterrows():
        model = row["Model"]
        authoritative = model != "BiLSTM"
        rows.append(
            {
                "Model": model,
                "Sensitivity source package": str(SENS_ROOT),
                "Hyperparameter or component tested": row["Hyperparameter comparison"],
                "Final setting": row["Final value"],
                "Alternatives tested": row["Alternative value"],
                "Evaluation basis": "Internal-validation nRMSE; delta relative to final setting",
                "Scope of analysis": "Compact representative local sensitivity; four tasks, folds 1 and 4, seed 42",
                "Number of evaluated fits/rows": 8,
                "Main result": row["Interpretation"],
                "Authoritative for manuscript?": "Yes" if authoritative else "No",
                "Notes": "Retained authoritative non-BiLSTM sensitivity evidence." if authoritative else "Superseded for affected targets by corrected true-BiLSTM local sensitivity.",
            }
        )
    for _, row in corr.iterrows():
        label = str(row["configuration"]).replace("_", " ")
        if row["configuration"] == "lower_capacity":
            tested = "Lower-capacity true BiLSTM: 64 units per direction in first layer and 64 second-layer units"
        elif row["configuration"] == "lower_learning_rate":
            tested = "Corrected final architecture with learning rate 0.0005"
        else:
            tested = "Corrected final architecture with learning rate 0.002"
        rows.append(
            {
                "Model": "BiLSTM",
                "Sensitivity source package": str(ROOT),
                "Hyperparameter or component tested": label,
                "Final setting": "Corrected true BiLSTM: 128 units per direction in first layer, second-layer units=128, learning rate=0.001",
                "Alternatives tested": tested,
                "Evaluation basis": "Internal-validation nRMSE; delta relative to corrected final BiLSTM",
                "Scope of analysis": "Corrected compact local sensitivity; four representative tasks, folds 1 and 4, seed 42",
                "Number of evaluated fits/rows": int(row["number_of_comparisons"]),
                "Main result": row["interpretation"],
                "Authoritative for manuscript?": "Yes",
                "Notes": "Authoritative BiLSTM sensitivity evidence after architecture correction.",
            }
        )
    for _, row in svr.iterrows():
        rows.append(
            {
                "Model": "SVR",
                "Sensitivity source package": str(SENS_ROOT),
                "Hyperparameter or component tested": f"{row['hyperparameter']}={row['value']}",
                "Final setting": "Task-specific C, epsilon, and gamma selected by internal-validation nRMSE from the approved RBF-SVR grid",
                "Alternatives tested": "C={1,3,10}; epsilon={0.05,0.10,0.20}; gamma multiplier={0.5,1,2} x gamma_scale",
                "Evaluation basis": "Internal-validation nRMSE only; timed-out or failed configurations retained and excluded from selection",
                "Scope of analysis": "Pixel-level pooled SVR validation grid; task-specific selection",
                "Number of evaluated fits/rows": int(row["candidate_count"]),
                "Main result": f"Median delta nRMSE relative to task-selected setting = {float(row['median_delta_nrmse_relative_to_task_selected']):.6f}",
                "Authoritative for manuscript?": "Yes",
                "Notes": "SVR is a pixel-level pooled conventional benchmark; settings may vary by target, land cover, and fold.",
            }
        )
    summary = pd.DataFrame(rows)
    summary_csv = REPORTS / "full_sensitivity_analysis_summary_table.csv"
    summary.to_csv(summary_csv, index=False)
    summary_docx = REPORTS / "full_sensitivity_analysis_summary_table.docx"
    csv_to_docx_table(summary_csv, summary_docx, "Full sensitivity analysis summary table")

    final_settings = {}
    for _, row in hparam.iterrows():
        final_settings[row["Model"]] = {
            "architecture": row["Architecture and main hyperparameters"],
            "training": row["Training configuration"],
            "selection": row["Selection basis"],
        }
    final_settings["BiLSTM"] = {
        "architecture": bilstm_hparam.iloc[0]["Corrected architecture and training setting"],
        "training": "AdamW; learning rate=0.001; weight decay=0.0; MSE loss; batch size=64; maximum epochs=8; patience=3; training sample cap=4000; gradient clipping max norm=1.0.",
        "selection": "Fixed corrected benchmark setting; validation used for early stopping only; held-out tests not used for hyperparameter selection.",
    }

    model_sections = []
    for model in ["LSTM", "BiLSTM", "TCN", "N-BEATS", "N-HiTS", "XGBoost", "SVR"]:
        auth_rows = summary[(summary["Model"] == model) & (summary["Authoritative for manuscript?"] == "Yes")]
        tested = "; ".join(auth_rows["Hyperparameter or component tested"].astype(str).tolist())
        results = "; ".join(auth_rows["Main result"].astype(str).tolist())
        setting = final_settings.get("RBF-epsilon-SVR" if model == "SVR" else model, final_settings.get(model, {}))
        classification = classify_model(model, auth_rows)
        model_sections.extend(
            [
                f"### {model}",
                "",
                f"- Final setting: {setting.get('architecture', 'Task-specific setting; see source tables.')}",
                f"- Alternatives/components tested: {tested}",
                "- Evaluation basis: internal-validation nRMSE only; held-out outer-test periods were not used.",
                f"- Scope: {auth_rows['Scope of analysis'].iloc[0] if len(auth_rows) else 'No authoritative rows found.'}",
                f"- Main conclusion: {classification}. {results}",
                "",
            ]
        )

    md = REPORTS / "full_sensitivity_analysis_status_report.md"
    now = datetime.now().isoformat(timespec="seconds")
    report_lines = [
        "# Full Hyperparameter-Sensitivity Analysis Status Report",
        "",
        f"Generated: {now}",
        "",
        "## 1. Overview",
        "",
        "This report consolidates the hyperparameter-sensitivity evidence completed so far for the forecasting models. The purpose was to check whether the final settings were reasonable and stable under local perturbations, using internal-validation data only. Held-out outer-test periods were not used for hyperparameter selection, model choice, or sensitivity interpretation.",
        "",
        "The sensitivity work should be read as representative and local, not as an exhaustive hyperparameter search. Final benchmark settings were not automatically changed by these diagnostics.",
        "",
        "## 2. Data sources and packages used",
        "",
        f"- Authoritative compact learned-model sensitivity package: `{SENS_ROOT}`",
        f"- Authoritative corrected BiLSTM rerun and sensitivity package: `{ROOT}`",
        f"- Authoritative SVR sensitivity/configuration evidence: `{SENS_ROOT / '05_svr'}` and `{SENS_ROOT / '06_tables'}`",
        f"- Superseded old BiLSTM rows exported as invalid/provisional: `{INVALID_OLD_BILSTM}` ({invalid_count} rows).",
        f"- Current hyperparameter audit table: `{FINAL_HPARAM}` plus corrected BiLSTM update `{CORR_BILSTM_HPARAM}`.",
        "",
        "## 3. Model-by-model summary",
        "",
        *model_sections,
        "## 4. SVR-specific clarification",
        "",
        "SVR is treated as a pixel-level pooled conventional machine-learning benchmark. One SVR model is fitted for each target, land-cover class, and rolling-origin fold using pooled pixel-time samples. Its sensitivity and selection focused on C, epsilon, and the RBF gamma multiplier. Selection was based on internal-validation nRMSE only, so the final SVR settings may be task-specific rather than one universal triplet.",
        "",
        f"Task-specific SVR settings are available in `{SENS_SVR_SELECTED}` and the final hyperparameter audit output.",
        "",
        "## 5. BiLSTM correction note",
        "",
        "An architecture-label inconsistency was identified in the earlier BiLSTM lineage: seven targets used a unidirectional first recurrent layer even though the model was labelled BiLSTM. Corrected true-BiLSTM benchmark reruns were completed for NDVI, GNDVI, NIRv, GOSIF/SIF, GPP, LST, and VPDmax. SPEI-1 and TWS were not rerun because their earlier BiLSTM implementation already used a bidirectional first layer.",
        "",
        "The corrected BiLSTM local sensitivity analysis in the corrected rerun package is now the authoritative BiLSTM sensitivity evidence. Old BiLSTM sensitivity rows for the seven affected targets should not be used in final figures, tables, or interpretation.",
        "",
        "## 6. Stability interpretation",
        "",
        "The retained non-BiLSTM sensitivity evidence is compact and representative: four target-land-cover tasks, folds 1 and 4, and seed 42. Several alternatives sometimes reduced validation nRMSE, but the diagnostics do not establish a globally superior replacement configuration. The corrected BiLSTM sensitivity is locally stable within the tested range: lower capacity and lower learning rate were generally worse, and higher learning rate was mixed rather than consistently better.",
        "",
        "## 7. Practical manuscript-use note",
        "",
        "This report is a consolidated evidence package for later manuscript and supplementary-material preparation. A final publication-ready supplementary figure and table can be prepared from these source tables after author approval and literature-supported wording. This report itself does not modify the manuscript.",
        "",
        "## 8. Summary table",
        "",
        summary.to_markdown(index=False),
        "",
        "## 9. Non-authoritative or superseded components",
        "",
        "- The original 1512-fit learned-model screening design was reduced before scientific interpretation because it was computationally disproportionate.",
        "- The full-confirmation stage was cancelled and should not be described as completed.",
        "- Old BiLSTM sensitivity rows for the seven affected targets are superseded by the corrected true-BiLSTM sensitivity analysis.",
        "- Existing non-BiLSTM and SVR sensitivity components are retained as authoritative.",
        "",
        "## 10. Non-modification confirmation",
        "",
        "No new sensitivity analysis was run by this reporting step. No forecasting model was rerun. No manuscript, supplementary Word file, accepted Figure 7, accepted Figure S6, or authoritative forecasting package was modified.",
        "",
    ]
    md.write_text("\n".join(report_lines), encoding="utf-8")
    docx = REPORTS / "full_sensitivity_analysis_status_report.docx"
    create_docx_from_markdown(md, docx)

    key_points = REPORTS / "sensitivity_analysis_key_points_for_author.md"
    key_points.write_text(
        "\n".join(
            [
                "# Sensitivity-analysis key points for author",
                "",
                "- Hyperparameter sensitivity was evaluated using internal-validation data only; outer tests were not used.",
                "- The learned-model sensitivity evidence is compact/local and representative, not exhaustive.",
                "- LSTM, TCN, N-BEATS, N-HiTS, and XGBoost use the retained compact balanced sensitivity outputs.",
                "- Corrected true-BiLSTM sensitivity replaces old affected BiLSTM sensitivity rows.",
                "- SVR is a pixel-level pooled conventional benchmark with task-specific internal-validation selection over C, epsilon, and RBF gamma multiplier.",
                "- No final hyperparameter was automatically replaced based on these diagnostics.",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    return md, docx, summary_csv, summary_docx, key_points


def classify_model(model: str, rows: pd.DataFrame) -> str:
    if model == "BiLSTM":
        return "Stable within the tested corrected local range"
    if model == "SVR":
        return "Task-dependent, with selection performed separately by internal-validation nRMSE"
    text = " ".join(rows["Main result"].astype(str).tolist()).lower()
    if "some local sensitivity" in text:
        return "Slightly sensitive but acceptable within the compact representative analysis"
    return "Stable within the tested range"


def write_patch_report(fig7_replaced: int, s6_monthly_replaced: int, s6_pixel_replaced: int, sensitivity_paths: tuple[Path, Path, Path, Path, Path]) -> Path:
    fig7_audit = pd.read_csv(DOWNSTREAM / "Figure7_BiLSTM_patch_audit.csv")
    s6_audit = pd.read_csv(DOWNSTREAM / "FigureS6_BiLSTM_patch_audit.csv")
    changed_fig7 = int(fig7_audit["changed"].sum())
    changed_s6 = int(s6_audit["changed"].sum())
    lines = [
        "# Final Figure Patch Report",
        "",
        f"Generated: {datetime.now().isoformat(timespec='seconds')}",
        "",
        "## Summary",
        "",
        "This patch recreated the accepted polar Figure 7 design and the accepted Supplementary Figure S6 design after replacing only the values affected by the corrected true-BiLSTM rerun. No forecasting model, hyperparameter search, validation routine, manuscript file, supplementary Word file, or accepted figure package was modified.",
        "",
        "## Accepted Figure 7 layout/source evidence",
        "",
        f"- Accepted polar plotting source table: `{FIG7_ACCEPTED_TABLE}`",
        f"- Accepted polar rebuild report: `{FIG7_ACCEPTED_REPORT}`",
        f"- Accepted polar v2 formatting report: `{FIG7_ACCEPTED_V2_REPORT}`",
        f"- Accepted polar v2 PDF: `{FIG7_ACCEPTED_V2_PDF}`",
        "- A standalone saved polar renderer script was not present in the accepted polar folder. The corrected renderer in this package recreates the documented accepted polar layout from the accepted plotting table and v2 formatting report.",
        "",
        "## Accepted Supplementary Figure S6 renderer/source evidence",
        "",
        f"- Accepted S6-with-SVR renderer: `{S6_ACCEPTED_RENDERER}`",
        f"- Accepted S6-with-SVR figure: `{S6_ACCEPTED_FIG}`",
        f"- Accepted combined monthly source: `{S6_ACCEPTED_MONTHLY}`",
        f"- Accepted combined pixel source: `{S6_ACCEPTED_PIXEL}`",
        f"- Accepted panel C timeline: `{S6_ACCEPTED_PANEL_C}`",
        "",
        "## Corrected BiLSTM source tables",
        "",
        f"- Corrected pixel metrics: `{CORRECTED_PIXEL}`",
        f"- Corrected monthly metrics: `{CORRECTED_MONTHLY}`",
        "",
        "## Patch audit",
        "",
        f"- Figure 7 BiLSTM plotting rows replaced: {fig7_replaced}.",
        f"- Figure 7 metric entries changed in audit table: {changed_fig7}.",
        f"- Figure S6 monthly BiLSTM source rows replaced: {s6_monthly_replaced}.",
        f"- Figure S6 pixel BiLSTM source rows replaced: {s6_pixel_replaced}.",
        f"- Figure S6 plotted summary entries changed in audit table: {changed_s6}.",
        "- SPEI-1 and TWS BiLSTM values were retained from the accepted sources.",
        "- All non-BiLSTM values were retained from the accepted sources.",
        "- SVR values were retained from the accepted S6-with-SVR source.",
        "- Panel (c) of S6 was copied from the accepted rolling-origin timeline source and did not change.",
        "",
        "## Output paths",
        "",
        f"- Final corrected polar Figure 7 PNG: `{DOWNSTREAM / 'final_corrected_Figure_7_polar.png'}`",
        f"- Final corrected polar Figure 7 PDF: `{DOWNSTREAM / 'final_corrected_Figure_7_polar.pdf'}`",
        f"- Final corrected Figure 7 source table: `{DOWNSTREAM / 'final_corrected_Figure_7_source_table.csv'}`",
        f"- Final corrected S6 PNG: `{DOWNSTREAM / 'final_corrected_Figure_S6.png'}`",
        f"- Final corrected S6 PDF: `{DOWNSTREAM / 'final_corrected_Figure_S6.pdf'}`",
        f"- Final corrected S6 panel (a) source: `{DOWNSTREAM / 'final_corrected_Figure_S6_panel_a_source.csv'}`",
        f"- Final corrected S6 panel (b) source: `{DOWNSTREAM / 'final_corrected_Figure_S6_panel_b_source.csv'}`",
        f"- Final corrected S6 panel (c) source: `{DOWNSTREAM / 'final_corrected_Figure_S6_panel_c_source.csv'}`",
        f"- Figure 7 patch audit: `{DOWNSTREAM / 'Figure7_BiLSTM_patch_audit.csv'}`",
        f"- Figure S6 patch audit: `{DOWNSTREAM / 'FigureS6_BiLSTM_patch_audit.csv'}`",
        f"- Figure 7 side-by-side comparison: `{DOWNSTREAM / 'Figure7_previous_vs_final_corrected_side_by_side.png'}`",
        f"- Figure S6 side-by-side comparison: `{DOWNSTREAM / 'FigureS6_previous_vs_final_corrected_side_by_side.png'}`",
        "",
        "## Sensitivity report outputs",
        "",
        f"- Full sensitivity status report MD: `{sensitivity_paths[0]}`",
        f"- Full sensitivity status report DOCX: `{sensitivity_paths[1]}`",
        f"- Full sensitivity summary table CSV: `{sensitivity_paths[2]}`",
        f"- Full sensitivity summary table DOCX: `{sensitivity_paths[3]}`",
        f"- Key points helper: `{sensitivity_paths[4]}`",
        "",
        "## Final confirmations",
        "",
        "- No model fitting was performed.",
        "- The polar Figure 7 design was preserved from the accepted polar source/design evidence.",
        "- The original S6 horizontal point/error-bar plus rolling-origin design was preserved.",
        "- SARIMA is absent from S6 and SVR remains in S6.",
        "- GNDVI and SPEI-1 retain their accepted no-detrending values except where GNDVI BiLSTM is specifically replaced by the corrected true-BiLSTM rerun.",
    ]
    path = REPORTS / "final_figure_patch_report.md"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def main() -> None:
    DOWNSTREAM.mkdir(parents=True, exist_ok=True)
    REPORTS.mkdir(parents=True, exist_ok=True)
    accepted_fig7, patched_fig7, fig7_replaced = patch_figure7_source()
    prev_fig7_png = DOWNSTREAM / "previous_accepted_Figure_7_polar_reconstructed.png"
    prev_fig7_pdf = DOWNSTREAM / "previous_accepted_Figure_7_polar_reconstructed.pdf"
    final_fig7_png = DOWNSTREAM / "final_corrected_Figure_7_polar.png"
    final_fig7_pdf = DOWNSTREAM / "final_corrected_Figure_7_polar.pdf"
    render_polar_figure(accepted_fig7, prev_fig7_png, prev_fig7_pdf)
    render_polar_figure(patched_fig7, final_fig7_png, final_fig7_pdf)

    monthly_new, pixel_new, panel_a, panel_b, panel_c, s6_monthly_replaced, s6_pixel_replaced = patch_s6_sources()
    prev_panel_a = pd.read_csv(S6_ACCEPTED_PANEL_A)
    prev_panel_b = pd.read_csv(S6_ACCEPTED_PANEL_B)
    prev_panel_c = pd.read_csv(S6_ACCEPTED_PANEL_C)
    prev_s6_png = DOWNSTREAM / "previous_accepted_Figure_S6_reconstructed.png"
    prev_s6_pdf = DOWNSTREAM / "previous_accepted_Figure_S6_reconstructed.pdf"
    final_s6_png = DOWNSTREAM / "final_corrected_Figure_S6.png"
    final_s6_pdf = DOWNSTREAM / "final_corrected_Figure_S6.pdf"
    render_s6(prev_panel_a, prev_panel_b, prev_panel_c, prev_s6_png, prev_s6_pdf)
    render_s6(panel_a, panel_b, panel_c, final_s6_png, final_s6_pdf)

    make_side_by_side(prev_fig7_png, final_fig7_png, DOWNSTREAM / "Figure7_previous_vs_final_corrected_side_by_side.png", "Accepted polar Figure 7 source", "Final corrected BiLSTM patch", max_height=2200)
    make_side_by_side(prev_s6_png, final_s6_png, DOWNSTREAM / "FigureS6_previous_vs_final_corrected_side_by_side.png", "Accepted S6 source", "Final corrected BiLSTM patch", max_height=1600)

    sensitivity_paths = build_sensitivity_status_report()
    report = write_patch_report(fig7_replaced, s6_monthly_replaced, s6_pixel_replaced, sensitivity_paths)

    status = {
        "generated": datetime.now().isoformat(timespec="seconds"),
        "no_model_rerun": True,
        "accepted_figure7_layout_evidence": str(FIG7_ACCEPTED_V2_REPORT),
        "accepted_figure7_renderer_note": "No standalone polar renderer script was found in the accepted polar package; accepted source table and v2 formatting report were used.",
        "accepted_s6_renderer": str(S6_ACCEPTED_RENDERER),
        "figure7_bilstm_plotting_rows_replaced": fig7_replaced,
        "figure_s6_bilstm_monthly_source_rows_replaced": s6_monthly_replaced,
        "figure_s6_bilstm_pixel_source_rows_replaced": s6_pixel_replaced,
        "outputs": {
            "figure7_png": str(final_fig7_png),
            "figure7_pdf": str(final_fig7_pdf),
            "figure_s6_png": str(final_s6_png),
            "figure_s6_pdf": str(final_s6_pdf),
            "patch_report": str(report),
            "sensitivity_report_md": str(sensitivity_paths[0]),
            "sensitivity_report_docx": str(sensitivity_paths[1]),
            "sensitivity_summary_csv": str(sensitivity_paths[2]),
            "sensitivity_summary_docx": str(sensitivity_paths[3]),
        },
        "hashes": {
            "figure7_source": sha256(DOWNSTREAM / "final_corrected_Figure_7_source_table.csv"),
            "s6_panel_a": sha256(DOWNSTREAM / "final_corrected_Figure_S6_panel_a_source.csv"),
            "s6_panel_b": sha256(DOWNSTREAM / "final_corrected_Figure_S6_panel_b_source.csv"),
            "s6_panel_c": sha256(DOWNSTREAM / "final_corrected_Figure_S6_panel_c_source.csv"),
        },
    }
    (REPORTS / "final_figure_patch_status.json").write_text(json.dumps(status, indent=2), encoding="utf-8")
    print(json.dumps(status, indent=2))


if __name__ == "__main__":
    main()
